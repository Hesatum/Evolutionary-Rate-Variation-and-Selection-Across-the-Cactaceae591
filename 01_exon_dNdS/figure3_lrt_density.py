#!/usr/bin/env python3
"""
figure3_lrt_density.py

Generates Figure 3: kernel-density plots of the likelihood-ratio test (LRT)
statistics (2Δℓ) for the site model (M1a vs. M2a, df=2) and the branch model
(M0 vs. Branch, df=6), with the χ² critical values at P=0.05/0.01 and a rug of
individual loci (manuscript Methods 2.3 / Figure 3 legend).

One exon has a branch-model 2Δℓ ≈ -49, an impossible value for nested models
that signals a codeml optimisation failure; it is excluded from the branch
panel's KDE and count (see Figure 3 legend) but not from the underlying data.

Expects an Excel workbook with one sheet per LRT comparison ("M1a vs M2a",
"M0 vs Branch"), each with a "2Δℓ" column (one row per exon).

Usage
-----
    python figure3_lrt_density.py --xlsx resultados_2026_frame_validated.xlsx \
        --site-sheet "M1a vs M2a" --branch-sheet "M0 vs Branch" \
        --out-prefix Figure3_LRT_density
"""

import argparse

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from matplotlib.ticker import MaxNLocator
from scipy import stats

C_LINE = '#08306b'   # density outline (dark blue)
C_FILL = '#eef2f7'   # full area (very light)
C_REJ = '#e69f00'    # rejection region (Okabe-Ito orange)
C_CRIT = '#b2182b'   # critical-value lines (dark red)
C_SIG = '#e69f00'    # significant rug ticks
C_NS = '#9e9e9e'     # non-significant rug ticks


def col(ws, name):
    hdr = [c.value for c in ws[1]]
    i = hdr.index(name)
    return [r[i] for r in ws.iter_rows(min_row=2, values_only=True)]


def lrt_of(wb, sheet, column):
    ws = wb[sheet]
    return np.array([v for v in col(ws, column) if isinstance(v, (int, float))], float)


def panel(ax, data, df, xlim, label, title, show_p01=True):
    crit05 = stats.chi2.ppf(0.95, df)
    crit01 = stats.chi2.ppf(0.99, df)
    n = data.size
    nsig = int((data > crit05).sum())
    on = data[(data >= xlim[0]) & (data <= xlim[1])]

    kde = stats.gaussian_kde(on)
    xs = np.linspace(xlim[0], xlim[1], 600)
    ys = kde(xs)

    ax.fill_between(xs, ys, color=C_FILL, zorder=1)
    rej = xs >= crit05
    ax.fill_between(xs[rej], ys[rej], color=C_REJ, alpha=0.55, lw=0, zorder=2)
    ax.plot(xs, ys, color=C_LINE, lw=1.3, zorder=4)

    ax.axvline(crit05, color=C_CRIT, lw=1.1, ls='--', zorder=5)
    if show_p01:
        ax.axvline(crit01, color=C_CRIT, lw=0.9, ls=':', zorder=5)

    ymax = ys.max()
    y0 = -ymax * 0.085
    ax.set_ylim(y0 * 1.7, ymax * 1.30)
    ax.set_xlim(*xlim)
    ax.axhline(0, color='#bbbbbb', lw=0.6, zorder=0)

    ax.text(crit05, ymax * 1.17, f'$\\chi^2_{{0.05}}$ = {crit05:.2f}',
            color=C_CRIT, fontsize=7.5, ha='center', va='bottom')
    if show_p01:
        ax.text(crit01, ymax * 1.02, f'$\\chi^2_{{0.01}}$ = {crit01:.2f}',
                color=C_CRIT, fontsize=7, ha='center', va='bottom', alpha=0.85)

    sig = on[on > crit05]
    ns = on[on <= crit05]
    ax.plot(ns, np.full_like(ns, y0), '|', color=C_NS, ms=5, mew=0.7, alpha=0.7, zorder=3)
    ax.plot(sig, np.full_like(sig, y0), '|', color=C_SIG, ms=6, mew=1.0, zorder=3)

    txt = f'$n$ = {n} exons\n{nsig} significant ($p$ < 0.05)'
    ax.text(0.97, 0.93, txt, transform=ax.transAxes, ha='right', va='top',
            fontsize=7.5, linespacing=1.4,
            bbox=dict(boxstyle='round,pad=0.4', fc='white', ec='#cccccc', lw=0.6))

    ax.set_title(title, fontsize=9.5, fontweight='bold', pad=6, loc='center')
    ax.set_ylabel('density')
    ax.yaxis.set_major_locator(MaxNLocator(4))
    ax.text(-0.085, 1.04, label, transform=ax.transAxes, fontsize=12,
            fontweight='bold', va='bottom', ha='left')


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                      formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--xlsx", required=True, help="Workbook with per-exon LRT values.")
    parser.add_argument("--site-sheet", default="M1a vs M2a")
    parser.add_argument("--branch-sheet", default="M0 vs Branch")
    parser.add_argument("--lrt-column", default="2Δℓ")
    parser.add_argument("--site-xlim", nargs=2, type=float, default=[-2, 80])
    parser.add_argument("--branch-xlim", nargs=2, type=float, default=[-2, 30])
    parser.add_argument("--out-prefix", default="Figure3_LRT_density")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    site = lrt_of(wb, args.site_sheet, args.lrt_column)          # df = 2
    branch = lrt_of(wb, args.branch_sheet, args.lrt_column)      # df = 6

    # Exclude branch exons with a negative 2Δℓ: a negative LRT for nested models
    # is impossible and signals a codeml optimisation failure.
    n_branch_neg = int((branch < 0).sum())
    branch = branch[branch >= 0]

    mpl.rcParams.update({
        'font.family': 'Arial', 'font.size': 9,
        'axes.linewidth': 0.8, 'axes.spines.top': False, 'axes.spines.right': False,
        'xtick.direction': 'out', 'ytick.direction': 'out',
        'xtick.major.size': 3, 'ytick.major.size': 3,
        'pdf.fonttype': 42, 'ps.fonttype': 42,
    })

    fig, (axa, axb) = plt.subplots(2, 1, figsize=(5.3, 6.4))
    panel(axa, site, df=2, xlim=tuple(args.site_xlim), label='A',
          title='Site model (M1a vs. M2a)')
    panel(axb, branch, df=6, xlim=tuple(args.branch_xlim), label='B',
          title='Branch model (M0 vs. Branch)')
    print(f'branch exons excluded (negative 2dℓ): {n_branch_neg}; branch n now {branch.size}')

    axb.set_xlabel('likelihood-ratio test statistic, 2Δℓ')
    axa.set_xlabel('likelihood-ratio test statistic, 2Δℓ')
    fig.subplots_adjust(left=0.12, right=0.97, top=0.94, bottom=0.08, hspace=0.42)

    for ext in ('pdf', 'png', 'tif'):
        fig.savefig(f'{args.out_prefix}.{ext}', dpi=600, bbox_inches='tight')
    print(f'saved {args.out_prefix}.pdf / .png / .tif')


if __name__ == '__main__':
    main()
